"""Excel workbook for the forward curve."""
from __future__ import annotations

import datetime as dt

from openpyxl import Workbook
from openpyxl.chart import LineChart, Reference
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .model import Curve

ARIAL = "Arial"
BLUE = "0000FF"
F_BASE = Font(name=ARIAL, size=10)
F_INPUT = Font(name=ARIAL, size=10, color=BLUE)
F_HEAD = Font(name=ARIAL, size=10, bold=True, color="FFFFFF")
F_TITLE = Font(name=ARIAL, size=14, bold=True)
F_BOLD = Font(name=ARIAL, size=10, bold=True)
F_NOTE = Font(name=ARIAL, size=9, italic=True, color="595959")
FILL_HEAD = PatternFill("solid", fgColor="1F3A5F")
FILL_INPUT = PatternFill("solid", fgColor="FFFF00")
FILL_EXTRAP = PatternFill("solid", fgColor="F2F2F2")
THIN = Side(style="thin", color="BFBFBF")
PCT = "0.000%"
DATE = "yyyy-mm-dd"

FIRST = 12  # first data row on the curve sheet


def _header(ws, row: int, labels: list[str], widths: list[int] | None = None):
    for c, label in enumerate(labels, 1):
        cell = ws.cell(row=row, column=c, value=label)
        cell.font, cell.fill = F_HEAD, FILL_HEAD
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        if widths:
            ws.column_dimensions[get_column_letter(c)].width = widths[c - 1]
    ws.row_dimensions[row].height = 30


def build_workbook(curve: Curve, fixings: dict[dt.date, float], base_url: str, path: str) -> None:
    wb = Workbook()
    n = len(curve.months)
    last = FIRST + n - 1

    # ---------------------------------------------------------------- Forward Curve
    ws = wb.active
    ws.title = "Forward Curve"
    ws["A1"] = "SOFR Forward Curve: monthly, 10 years"
    ws["A1"].font = F_TITLE
    rows = [
        ("Curve date (CME settlement)", dt.date.fromisoformat(curve.as_of.isoformat()), DATE),
        ("Last SOFR fixing", curve.last_fixing[1] / 100, PCT),
        ("Generated", dt.datetime.utcnow().strftime("%Y-%m-%d %H:%M UTC"), None),
    ]
    for r, (label, val, fmt) in enumerate(rows, 2):
        ws.cell(row=r, column=1, value=label).font = F_BOLD
        c = ws.cell(row=r, column=2, value=val)
        c.font = F_BASE
        if fmt:
            c.number_format = fmt
    ws["C3"] = f"on {curve.last_fixing[0].isoformat()} (NY Fed)"
    ws["C3"].font = F_NOTE

    ws["A6"] = "Loan inputs (edit the yellow cells)"
    ws["A6"].font = F_BOLD
    inputs = [
        ("Spread over SOFR (bp)", 250, "0"),
        ("SOFR floor", 0.0, PCT),
        ("SOFR cap strike (blank = no cap)", None, PCT),
    ]
    for r, (label, val, fmt) in enumerate(inputs, 7):
        ws.cell(row=r, column=1, value=label).font = F_BASE
        c = ws.cell(row=r, column=2, value=val)
        c.font, c.fill, c.number_format = F_INPUT, FILL_INPUT, fmt
    ws["C7"] = "Example: 250 = SOFR + 2.50%"
    ws["C9"] = "Example: 4.00% caps the index at 4.00%"
    for a in ("C7", "C9"):
        ws[a].font = F_NOTE

    _header(ws, FIRST - 1, [
        "Reset Date", "End Date", "Days", "1M Fwd SOFR", "Avg Overnight SOFR",
        "Futures-Implied (no convexity)", "Convexity (bp)", "Status",
        "Index Rate (after floor / cap)", "All-in Rate",
    ], [13, 13, 7, 13, 13, 15, 11, 15, 15, 12])

    for k, m in enumerate(curve.months):
        r = FIRST + k
        vals = [
            dt.date.fromisoformat(m["reset_date"]), dt.date.fromisoformat(m["end_date"]), m["days"],
            m["forward_1m"] / 100, m["avg_overnight"] / 100, m["futures_implied"] / 100,
            m["convexity_bp"], m["status"].replace("_", " "),
        ]
        fmts = [DATE, DATE, "0", PCT, PCT, PCT, "0.0", None]
        for c, (v, f) in enumerate(zip(vals, fmts), 1):
            cell = ws.cell(row=r, column=c, value=v)
            cell.font = F_INPUT if c in (4, 5, 6, 7) else F_BASE
            if f:
                cell.number_format = f
        ws.cell(row=r, column=9, value=f'=MAX($B$8,IF($B$9="",D{r},MIN(D{r},$B$9)))').number_format = PCT
        ws.cell(row=r, column=10, value=f"=I{r}+$B$7/10000").number_format = PCT
        for c in (9, 10):
            ws.cell(row=r, column=c).font = F_BASE
        if m["status"] == "extrapolated":
            for c in range(1, 11):
                ws.cell(row=r, column=c).fill = FILL_EXTRAP
        for c in range(1, 11):
            ws.cell(row=r, column=c).border = Border(bottom=THIN)
    ws.freeze_panes = ws.cell(row=FIRST, column=2)
    note_r = last + 2
    notes = [
        "1M Fwd SOFR: compounded SOFR over the calendar month, approximately the 1M Term SOFR fixing on the reset date.",
        "Status: 'partially fixed' months blend published SOFR fixings with forwards; 'extrapolated' months (grey) lie "
        "beyond the last listed SR3 contract and hold its level flat.",
        "Blue numbers are model outputs pasted as values; black cells are formulas driven by the yellow inputs.",
        f"Live data: {base_url}/api/v1/latest.csv  (Excel: Data > From Web, paste the URL, then Refresh All).",
    ]
    for i, t in enumerate(notes):
        ws.cell(row=note_r + i, column=1, value=t).font = F_NOTE

    chart = LineChart()
    chart.title = f"SOFR forward curve, {curve.as_of.isoformat()}"
    chart.y_axis.title = "Rate"
    chart.y_axis.number_format = "0.00%"
    chart.x_axis.number_format = "mmm-yy"
    chart.x_axis.title = None
    chart.height, chart.width = 9, 22
    data = Reference(ws, min_col=4, max_col=4, min_row=FIRST - 1, max_row=last)
    chart.add_data(data, titles_from_data=True)
    data2 = Reference(ws, min_col=6, max_col=6, min_row=FIRST - 1, max_row=last)
    chart.add_data(data2, titles_from_data=True)
    chart.set_categories(Reference(ws, min_col=1, min_row=FIRST, max_row=last))
    chart.series[0].graphicalProperties.line.solidFill = "1F3A5F"
    chart.series[0].graphicalProperties.line.width = 28000
    chart.series[1].graphicalProperties.line.solidFill = "A6A6A6"
    chart.series[1].graphicalProperties.line.dashStyle = "dash"
    for s in chart.series:
        s.smooth = False
        s.marker.symbol = "none"
    chart.x_axis.tickLblSkip = 12
    chart.x_axis.delete = False
    chart.y_axis.delete = False
    ws.add_chart(chart, "L2")

    # ---------------------------------------------------------------- Annual
    wa = wb.create_sheet("Annual")
    wa["A1"] = "Calendar-year averages (simple average of monthly values)"
    wa["A1"].font = F_BOLD
    _header(wa, 3, ["Year", "Months in Curve", "Avg 1M Fwd SOFR", "Avg All-in Rate"], [10, 16, 18, 18])
    rng = lambda col: f"'Forward Curve'!${col}${FIRST}:${col}${last}"  # noqa: E731
    years = sorted({int(m["month"][:4]) for m in curve.months})
    for k, y in enumerate(years):
        r = 4 + k
        wa.cell(row=r, column=1, value=y).number_format = "0"
        crit = f'{rng("A")},">="&DATE(A{r},1,1),{rng("A")},"<"&DATE(A{r}+1,1,1)'
        wa.cell(row=r, column=2, value=f"=COUNTIFS({crit})")
        wa.cell(row=r, column=3, value=f'=IFERROR(AVERAGEIFS({rng("D")},{crit}),"")').number_format = PCT
        wa.cell(row=r, column=4, value=f'=IFERROR(AVERAGEIFS({rng("J")},{crit}),"")').number_format = PCT
        for c in range(1, 5):
            wa.cell(row=r, column=c).font = F_INPUT if c == 1 else F_BASE
            wa.cell(row=r, column=c).border = Border(bottom=THIN)
    wa.cell(row=5 + len(years), column=1,
            value="Years with fewer than 12 months are partial (curve starts in the current month).").font = F_NOTE

    # ---------------------------------------------------------------- Futures Inputs
    wf = wb.create_sheet("Futures Inputs")
    wf["A1"] = f"CME SOFR futures settlements used in the fit (trade date {curve.as_of.isoformat()})"
    wf["A1"].font = F_BOLD
    cols = ["Contract", "Reference Start", "Reference End", "Settle", "Implied Rate",
            "Convexity (bp)", "Model Rate", "Residual (bp)", "Open Interest"]
    _header(wf, 3, cols, [12, 15, 15, 10, 12, 12, 12, 12, 14])
    for k, i in enumerate(curve.instruments):
        r = 4 + k
        vals = [i["contract"], dt.date.fromisoformat(i["ref_start"]), dt.date.fromisoformat(i["ref_end"]),
                i["settle"], None, i["convexity_bp"], i["model_rate"] / 100, i["residual_bp"], i["open_interest"]]
        fmts = [None, DATE, DATE, "0.0000", PCT, "0.00", PCT, "0.00", "#,##0"]
        for c, (v, f) in enumerate(zip(vals, fmts), 1):
            cell = wf.cell(row=r, column=c, value=v)
            cell.font = F_INPUT if c in (4, 6, 7, 8, 9) else F_BASE
            if f:
                cell.number_format = f
        wf.cell(row=r, column=5, value=f"=(100-D{r})/100").font = F_BASE
        for c in range(1, 10):
            wf.cell(row=r, column=c).border = Border(bottom=THIN)
    wf.freeze_panes = "B4"
    s = curve.stats
    wf.cell(row=5 + len(curve.instruments), column=1, value=(
        f"Fit: {s['instruments']} contracts, RMS residual {s['rms_residual_bp']} bp, "
        f"max {s['max_abs_residual_bp']} bp. Residual = model futures rate - market rate. "
        "Source: CME Group daily settlements (SR3, SR1).")).font = F_NOTE

    # ---------------------------------------------------------------- SOFR Fixings
    wx = wb.create_sheet("SOFR Fixings")
    wx["A1"] = "Published SOFR used for elapsed days (source: Federal Reserve Bank of New York)"
    wx["A1"].font = F_BOLD
    _header(wx, 3, ["Date", "SOFR"], [13, 10])
    for k, (d, v) in enumerate(sorted(fixings.items(), reverse=True)):
        a = wx.cell(row=4 + k, column=1, value=d)
        b = wx.cell(row=4 + k, column=2, value=v / 100)
        a.number_format, b.number_format = DATE, PCT
        a.font, b.font = F_BASE, F_INPUT

    # ---------------------------------------------------------------- Methodology
    wm = wb.create_sheet("Methodology")
    wm.column_dimensions["A"].width = 120
    p = curve.params
    lines = [
        ("Methodology", F_TITLE),
        ("Inputs: CME Three-Month (SR3) and One-Month (SR1) SOFR futures daily settlements; NY Fed SOFR fixings.", F_BASE),
        ("SR1 settles on the calendar-day average of SOFR in its month; SR3 on business-day compounded SOFR from the "
         "3rd Wednesday of its named month to the 3rd Wednesday three months later.", F_BASE),
        ("Front end: overnight SOFR is constant between FOMC decision dates (rate changes effective the day after), "
         "the same step structure CME uses for Term SOFR. Beyond the last scheduled FOMC meeting: one level per month.", F_BASE),
        ("Levels are fitted by least squares to all SR3 quarterlies and to serial SR3 / SR1 contracts with open interest, "
         "with a light curvature penalty so the back end is smooth.", F_BASE),
        (f"Convexity: futures rates are reduced by a Hull-White one-factor adjustment (sigma = {p.hw_sigma:.2%} normal, "
         f"mean reversion a = {p.hw_a}). On this curve: {curve.months[36]['convexity_bp']:.0f} bp at 3 years, "
         f"{curve.months[60]['convexity_bp']:.0f} bp at 5 years, {curve.months[-1]['convexity_bp']:.0f} bp at 10 years. "
         "The 'Futures-Implied' column shows the curve with no adjustment.", F_BASE),
        ("Contracts beyond about 7 years carry little or no open interest; CME sets their settlements from spreads, so the "
         "last few years of the curve are indicative. Months after the last contract hold its level flat.", F_BASE),
        ("1M Fwd SOFR is compounded SOFR over each calendar month and approximates the 1M CME Term SOFR fixing on the "
         "reset date to within a few bp; it is not CME Term SOFR.", F_BASE),
        ("Not investment advice. Market data (c) CME Group and the Federal Reserve Bank of New York.", F_NOTE),
        ("", F_BASE),
        ("API (CORS-enabled, static JSON/CSV)", F_BOLD),
        (f"{base_url}/api/v1/latest.json", F_BASE),
        (f"{base_url}/api/v1/latest.csv", F_BASE),
        (f"{base_url}/api/v1/history/index.json  (list of dates; then /api/v1/history/YYYY-MM-DD.json or .csv)", F_BASE),
        (f"{base_url}/downloads/sofr-forward-curve.xlsx", F_BASE),
    ]
    for r, (t, f) in enumerate(lines, 1):
        c = wm.cell(row=r, column=1, value=t)
        c.font = f
        c.alignment = Alignment(wrap_text=True, vertical="top")

    wb.calculation.fullCalcOnLoad = True
    wb.save(path)
